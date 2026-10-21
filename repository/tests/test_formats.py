import hashlib
import io
import os
import posixpath
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import numpy as np
import pytest
from docx import Document as WordDocument
from docx.shared import Inches
from PIL import Image
from pypdf import PdfReader

from clearscore.core import Cancelled, CleanupError
from clearscore.formats import prepare_document, save_document
from tests.helpers import score_image


def make_docx(path, image_path):
    doc = WordDocument()
    doc.add_heading("Practice / Практика / مەشىق", 0)
    doc.add_paragraph("Keep this text exactly. 1 2 3 — Am C G7")
    doc.add_picture(str(image_path), width=Inches(5.1))
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Tempo = 62"
    table.cell(1, 1).text = "Capo 2"
    doc.sections[0].header.paragraphs[0].text = "Header must stay"
    doc.sections[0].footer.paragraphs[0].text = "Footer must stay"
    doc.save(path)
    return path


@pytest.mark.parametrize("suffix,format", [
    ("png", "PNG"), ("jpg", "JPEG"), ("jpeg", "JPEG"), ("bmp", "BMP"),
    ("webp", "WEBP"), ("tiff", "TIFF"), ("gif", "GIF"), ("avif", "AVIF"),
    ("jp2", "JPEG2000"), ("tga", "TGA"), ("ppm", "PPM"),
])
def test_image_decoders_and_lossless_output(tmp_path, suffix, format):
    path = tmp_path / ("input." + suffix)
    score_image().save(path, format)
    before_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    prepared = prepare_document(path)
    assert prepared.kind == "image"
    assert prepared.report.changed_pixels > 1000
    output = save_document(prepared, tmp_path / "result.png", path)
    decoded = Image.open(output)
    a, b = np.array(prepared.frames[0].original.convert("RGB")), np.array(decoded.convert("RGB"))
    cutoff = prepared.report.images[0].threshold
    assert np.array_equal(a[a.min(axis=2) < cutoff], b[a.min(axis=2) < cutoff])
    assert np.array_equal(np.array(decoded), np.array(prepared.frames[0].cleaned))
    assert before_hash == hashlib.sha256(path.read_bytes()).hexdigest()


def test_heic_decoder_and_lossless_export(tmp_path):
    source = Path(__file__).parent / "fixtures" / "synthetic-score.heic"
    doc = prepare_document(source)
    assert doc.report.changed_pixels > 1000
    target = save_document(doc, tmp_path / "decoded.png", source)
    assert np.array_equal(np.array(Image.open(target)), np.array(doc.frames[0].cleaned))


def test_alpha_is_kept_and_pdf_uses_lossless_samples(tmp_path):
    image = score_image().convert("RGBA")
    image.paste((33, 66, 99, 0), (0, 0, 80, 80))
    image.paste((229, 229, 229, 100), (100, 100, 150, 150))
    path = tmp_path / "transparent.png"
    image.save(path, dpi=(300, 300))
    doc = prepare_document(path)
    original, cleaned = np.array(image), np.array(doc.frames[0].cleaned)
    assert np.array_equal(original[:, :, 3], cleaned[:, :, 3])
    assert np.array_equal(original[original[:, :, 3] != 255], cleaned[original[:, :, 3] != 255])
    target = save_document(doc, tmp_path / "score.pdf", path)
    result = PdfReader(target)
    extracted = result.pages[0].images[0]
    assert extracted.indirect_reference.get_object()["/Filter"] == "/FlateDecode"
    assert np.array_equal(np.array(extracted.image.convert("RGBA")), cleaned)
    assert abs(float(result.pages[0].mediabox.width) - 192) < 0.01


def test_exif_orientation_is_applied_once(tmp_path):
    path = tmp_path / "rotated.jpg"
    exif = Image.Exif()
    exif[274] = 6
    score_image().save(path, exif=exif)
    doc = prepare_document(path)
    assert doc.frames[0].original.size == (1100, 800)
    output = save_document(doc, tmp_path / "clean.png", path)
    result = Image.open(output)
    assert result.size == (1100, 800)
    assert result.getexif().get(274, 1) == 1


@pytest.mark.parametrize("format,suffix", [("TIFF", "tiff"), ("GIF", "gif"), ("PNG", "png"), ("WEBP", "webp")])
def test_multiple_frames_never_silently_drop_pages(tmp_path, format, suffix):
    first, second = score_image(), score_image().rotate(180)
    path = tmp_path / ("pages." + suffix)
    first.save(path, format, save_all=True, append_images=[second], duration=400)
    doc = prepare_document(path)
    assert len(doc.frames) == 2
    assert doc.default_suffix == ".tiff"
    with pytest.raises(CleanupError) as failure:
        save_document(doc, tmp_path / "wrong.png", path)
    assert failure.value.message.key == "multi_export"
    output = save_document(doc, tmp_path / "all.tiff", path)
    with Image.open(output) as result:
        assert result.n_frames == 2
        for i in range(2):
            result.seek(i)
            assert np.array_equal(np.array(result), np.array(doc.frames[i].cleaned))
    pdf = save_document(doc, tmp_path / "all.pdf", path)
    assert len(PdfReader(pdf).pages) == 2


def test_docx_preserves_text_layout_and_relationships(tmp_path):
    image = tmp_path / "scan.jpg"
    score_image().save(image, quality=95)
    source = make_docx(tmp_path / "original.docx", image)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    doc = prepare_document(source)
    assert doc.kind == "docx"
    assert doc.report.changed_pixels > 1000
    target = save_document(doc, tmp_path / "result.docx", source)
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(target) as after:
        for name in before.namelist():
            if name.startswith("word/media/") or name.endswith(".rels") or name == "[Content_Types].xml":
                continue
            assert before.read(name) == after.read(name), name
        assert "word/media/image1.jpg" not in after.namelist()
        new_media = [name for name in after.namelist() if name.startswith("word/media/")]
        assert len(new_media) == 1 and new_media[0].endswith(".png")
        clean = Image.open(io.BytesIO(after.read(new_media[0])))
        assert np.array_equal(np.array(clean), np.array(doc.frames[0].cleaned))
        relationships = ET.fromstring(after.read("word/_rels/document.xml.rels"))
        for relationship in relationships:
            if relationship.attrib["Type"].endswith("/image"):
                assert posixpath.normpath("word/" + relationship.attrib["Target"]) in after.namelist()
        content_types = ET.fromstring(after.read("[Content_Types].xml"))
        assert any(node.get("PartName") == "/" + new_media[0] and node.get("ContentType") == "image/png" for node in content_types)
    # An independent Word library must still read the edited package successfully.
    result = WordDocument(target)
    assert result.paragraphs[1].text == "Keep this text exactly. 1 2 3 — Am C G7"
    assert len(result.inline_shapes) == 1
    assert result.inline_shapes[0].width == WordDocument(source).inline_shapes[0].width
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest


def test_source_overwrite_hardlink_and_cancellation(tmp_path):
    path = tmp_path / "score.png"
    score_image().save(path)
    doc = prepare_document(path)
    linked = tmp_path / "linked.png"
    os.link(path, linked)
    for output in (path, linked):
        with pytest.raises(CleanupError, match="原始"):
            save_document(doc, output, path)
    with pytest.raises(Cancelled):
        prepare_document(path, cancelled=lambda: True)
    assert not list(tmp_path.glob(".clearscore-*"))


def test_legacy_and_malformed_word_are_explained(tmp_path):
    with pytest.raises(CleanupError) as old:
        prepare_document(tmp_path / "old.doc")
    assert old.value.message.key == "legacy_word"
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"Not a ZIP")
    with pytest.raises(CleanupError) as invalid:
        prepare_document(bad)
    assert invalid.value.message.key == "process_failed"
    no_images = tmp_path / "text.docx"
    WordDocument().save(no_images)
    with pytest.raises(CleanupError) as empty:
        prepare_document(no_images)
    assert empty.value.message.key == "no_images"


def test_signed_docx_is_never_modified(tmp_path):
    image = tmp_path / "scan.png"
    score_image().save(image)
    path = make_docx(tmp_path / "signed.docx", image)
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("_xmlsignatures/sig1.xml", "<Signature/>")
    with pytest.raises(CleanupError) as error:
        prepare_document(path)
    assert error.value.message.key == "signed_docx"
